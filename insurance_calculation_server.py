#!/usr/bin/env python3
"""Local calculation and import service for premium-financing proposals.

The browser never attempts to reproduce insurer workbook formula chains.  It sends
the selected product inputs to this localhost service, which changes a disposable
copy of the approved proposal workbook and asks LibreOffice Calc to recalculate it.
Only the resulting quotation values are returned to the browser.
"""

from __future__ import annotations

import argparse
import cgi
import copy
import hashlib
import hmac
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("INSURANCE_DATA_DIR", str(ROOT / "data"))).resolve()
SOURCES_DIR = DATA_DIR / "product_sources"
MANIFEST_PATH = DATA_DIR / "products.json"
PUBLISHED_QUOTES_DIR = DATA_DIR / "published_quotes"
SOFFICE = os.environ.get("SOFFICE_PATH", "/opt/homebrew/bin/soffice")
CALCULATION_TIMEOUT_SECONDS = int(os.environ.get("CALCULATION_TIMEOUT_SECONDS", "150"))
ADMIN_UPLOAD_TOKEN = os.environ.get("ADMIN_UPLOAD_TOKEN", "")
CORS_ALLOWED_ORIGINS = {
    origin.strip().rstrip("/")
    for origin in os.environ.get(
        "CORS_ALLOWED_ORIGINS",
        "https://anzuanzu.github.io,http://localhost:8765,http://127.0.0.1:8765",
    ).split(",")
    if origin.strip()
}
MAX_UPLOAD_BYTES = 40 * 1024 * 1024
MAX_DECLARED_RATE = 0.20
MAX_QUOTE_YEARS = 100
MAX_SHARED_QUOTES_PER_PRODUCT = 160
QUOTE_DATA_GLOBAL_NAME = "INSURANCE_PRODUCT_QUOTE_DATA_BY_PROFILE"
DIVIDEND_OPTION_CONTINUE = "第7年起持續增購保額"
DIVIDEND_OPTION_VALUES = {
    "第7年起儲存生息",
    DIVIDEND_OPTION_CONTINUE,
    "當年度給付",
}
BUILT_IN_SOURCE_CODES = {"ABA", "WUS", "WEF", "WQ6", "WJA"}
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL)


class CalculationError(ValueError):
    """A proposal cannot be safely imported or calculated."""


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def quote_cache_key(quote: dict[str, Any]) -> str:
    """Return a stable, source-version-independent browser lookup key."""
    return "|".join((
        str(int(quote["age"])),
        "female" if quote.get("gender") == "female" else "male",
        format(float(quote["faceAmount"]), ".4f"),
        format(float(quote["declaredRate"]), ".4f"),
    ))


def quote_record_path(profile_id: str, quote: dict[str, Any]) -> Path:
    """Return the independent shared-record path for an exact quote.

    A quote is stored separately from the product manifest.  This lets Cloud
    Run instances publish different calculations concurrently without one
    instance overwriting another instance's in-memory manifest revision.
    """
    digest = hashlib.sha256(quote_cache_key(quote).encode("utf-8")).hexdigest()
    return PUBLISHED_QUOTES_DIR / profile_id / f"{digest}.json"


def disk_cached_quotes(profile: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Read independently published quotes for the current source workbook."""
    directory = PUBLISHED_QUOTES_DIR / profile["id"]
    if not directory.exists():
        return {}
    quotes: dict[str, dict[str, Any]] = {}
    try:
        paths = list(directory.glob("*.json"))
    except OSError:
        return {}
    for path in paths:
        try:
            record = json.loads(path.read_text("utf-8"))
            quote = record.get("quote") if isinstance(record, dict) else None
            if not isinstance(quote, dict) or quote.get("sourceHash") != profile.get("sourceHash"):
                continue
            quotes[quote_cache_key(quote)] = quote
        except (OSError, json.JSONDecodeError, KeyError, TypeError, ValueError):
            # A concurrent publisher can briefly expose an incomplete FUSE
            # directory listing.  Ignore that one record and use the next
            # request to pick it up once published.
            continue
    return quotes


def quote_data_payload(profile: dict[str, Any]) -> dict[str, Any]:
    """Public, exact Calc results that the browser may use without a quote call."""
    cache = profile.get("quoteCache") or {}
    quotes_by_key = {
        quote_cache_key(item["quote"]): item["quote"]
        for item in cache.values()
        if isinstance(item, dict) and isinstance(item.get("quote"), dict)
        and item["quote"].get("sourceHash") == profile.get("sourceHash")
    }
    # Disk entries take precedence because they are the latest exact results
    # published by any Cloud Run instance.
    quotes_by_key.update(disk_cached_quotes(profile))
    quotes = list(quotes_by_key.values())
    quotes.sort(key=lambda item: (item.get("age", 0), item.get("gender", ""), item.get("faceAmount", 0), item.get("declaredRate", 0)))
    return {
        "schemaVersion": 1,
        "profileId": profile["id"],
        "webCode": profile["webCode"],
        "sourceHash": profile["sourceHash"],
        "productVersion": profile.get("version", 1),
        "generatedAt": utc_now(),
        "quotes": quotes,
    }


def quote_data_version(profile: dict[str, Any]) -> str:
    payload = quote_data_payload(profile)
    # generatedAt is informational and must not invalidate an unchanged browser cache.
    payload.pop("generatedAt", None)
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(serialized).hexdigest()[:16]


def quote_data_javascript(profile: dict[str, Any]) -> str:
    payload = quote_data_payload(profile)
    return (
        "// Generated from the approved proposal workbook by LibreOffice Calc.\n"
        f"window.{QUOTE_DATA_GLOBAL_NAME} = window.{QUOTE_DATA_GLOBAL_NAME} || Object.create(null);\n"
        f"window.{QUOTE_DATA_GLOBAL_NAME}[{json.dumps(profile['id'], ensure_ascii=False)}] = Object.freeze("
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ");\n"
    )


def normalized_text(value: Any) -> str:
    return "".join(unicodedata.normalize("NFKC", str(value or "")).split()).casefold()


def has_financing_filename(filename: str) -> bool:
    return "保費融資" in unicodedata.normalize("NFKC", filename or "")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def as_number(value: Any, label: str) -> float:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        raise CalculationError(f"{label} 缺少可用數值。") from None
    if not numeric == numeric or numeric in (float("inf"), float("-inf")):
        raise CalculationError(f"{label} 不是有效數值。")
    return numeric


def currency_from_value(value: Any) -> str:
    text = str(value or "").upper()
    mapping = {"1": "TWD", "2": "USD", "3": "AUD", "4": "EUR", "5": "GBP"}
    if text in mapping:
        return mapping[text]
    for currency in ("USD", "TWD", "AUD", "EUR", "GBP"):
        if currency in text:
            return currency
    raise CalculationError("無法辨識商品幣別。")


def safe_filename(filename: str) -> str:
    name = Path(filename or "").name
    if not name or name in {".", ".."}:
        raise CalculationError("檔名無效。")
    return name


def name_destination(workbook: openpyxl.Workbook, name: str) -> tuple[str, str]:
    defined = workbook.defined_names.get(name)
    if not defined:
        raise CalculationError(f"建議書缺少命名欄位「{name}」。")
    try:
        sheet, cell_range = next(defined.destinations)
    except (StopIteration, AttributeError, ValueError) as exc:
        raise CalculationError(f"命名欄位「{name}」不是有效儲存格。") from exc
    if ":" in cell_range:
        raise CalculationError(f"命名欄位「{name}」必須是單一儲存格。")
    return sheet, cell_range.replace("$", "")


def optional_destination(workbook: openpyxl.Workbook, name: str) -> tuple[str, str] | None:
    if name not in workbook.defined_names:
        return None
    try:
        return name_destination(workbook, name)
    except CalculationError:
        return None


def read_named_value(workbook: openpyxl.Workbook, name: str, required: bool = True) -> Any:
    destination = optional_destination(workbook, name)
    if not destination:
        if required:
            raise CalculationError(f"建議書缺少命名欄位「{name}」。")
        return None
    sheet, cell = destination
    return workbook[sheet][cell].value


def read_named_range(workbook: openpyxl.Workbook, name: str) -> list[Any]:
    defined = workbook.defined_names.get(name)
    if not defined:
        raise CalculationError(f"建議書缺少結果欄位「{name}」。")
    try:
        sheet, cell_range = next(defined.destinations)
    except (StopIteration, AttributeError, ValueError) as exc:
        raise CalculationError(f"結果欄位「{name}」無法讀取。") from exc
    min_col, min_row, max_col, max_row = range_boundaries(cell_range.replace("$", ""))
    values: list[Any] = []
    worksheet = workbook[sheet]
    for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True):
        values.extend(row)
    return values


def read_standard_result_column(workbook: openpyxl.Workbook, column: str) -> list[Any]:
    """Read the standard financial-analysis table used by the approved templates.

    Workbook labels vary slightly (notably WUS), whereas the table positions are
    stable: O is total policy cash value and R is total general protection.
    """
    worksheet = workbook["利益分析表_基本"]
    return [worksheet[f"{column}{row}"].value for row in range(5, 5 + MAX_QUOTE_YEARS)]


def numeric_year_values(values: list[Any], label: str) -> list[float]:
    numeric: list[float] = []
    for year, value in enumerate(values[:MAX_QUOTE_YEARS], start=1):
        if value is None or value == "" or value in {"-", "—", "–"}:
            numeric.append(None)  # type: ignore[arg-type]
            continue
        if isinstance(value, str) and value.startswith("#"):
            raise CalculationError(f"{label}第 {year} 年出現公式錯誤：{value}")
        try:
            numeric.append(float(value))
        except (TypeError, ValueError):
            # A proposal may mark post-coverage years with text such as「無」.
            # It is not a calculation failure and must remain unavailable, not zero.
            numeric.append(None)  # type: ignore[arg-type]
    if not any(value is not None for value in numeric):
        raise CalculationError(f"{label}沒有可用結果。")
    return numeric


def short_product_name(name: str, code: str) -> str:
    compact = re.sub(r"^國泰人壽", "", name).replace("利率變動型", "")
    compact = compact.replace("美元終身壽險", "").replace("終身壽險", "")
    compact = compact.replace("(定期給付型)", "").strip()
    return compact[:28] or code


def workbook_sheet_paths(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        item.attrib.get("Id"): item.attrib.get("Target", "")
        for item in relationships.findall(f"{{{NS_PACKAGE_REL}}}Relationship")
    }
    paths: dict[str, str] = {}
    for sheet in workbook_xml.findall(f".//{{{NS_MAIN}}}sheet"):
        relation_id = sheet.attrib.get(f"{{{NS_REL}}}id")
        target = targets.get(relation_id, "")
        if target:
            paths[sheet.attrib["name"]] = "xl/" + target.lstrip("/")
    return paths


def find_dividend_option_input(path: Path, values_workbook: openpyxl.Workbook) -> tuple[tuple[str, str], str]:
    """Find the cell which drives the year-seven dividend handling.

    The approved templates expose this choice through the ``預算`` mapping
    formula.  We follow its direct sheet/cell reference instead of scanning for
    a matching display label, because a proposal can contain several example
    choices that are not the active calculation input.
    """
    formulas_workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    reference_pattern = re.compile(r"(?:^|[=(,])(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fff ]+))!\$?([A-Z]{1,3})\$?(\d+)")
    for worksheet in formulas_workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                formula = cell.value
                if not isinstance(formula, str) or not formula.startswith("=") or DIVIDEND_OPTION_CONTINUE not in formula:
                    continue
                for quoted_sheet, plain_sheet, column, row_number in reference_pattern.findall(formula):
                    sheet_name = quoted_sheet or plain_sheet
                    cell_ref = f"{column}{row_number}"
                    if sheet_name not in values_workbook.sheetnames:
                        continue
                    active_value = values_workbook[sheet_name][cell_ref].value
                    if isinstance(active_value, str) and active_value.strip() in DIVIDEND_OPTION_VALUES:
                        return (sheet_name, cell_ref), active_value.strip()
    raise CalculationError("無法定位增值回饋分享金給付方式輸入格；此建議書不能保證採用第7年起持續增購保額。")


def extract_profile(path: Path, original_filename: str | None = None) -> dict[str, Any]:
    filename = safe_filename(original_filename or path.name)
    if not has_financing_filename(filename):
        raise CalculationError("只接受檔名包含「保費融資」的建議書。")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise CalculationError("只接受 .xlsx 或 .xlsm 建議書。")
    try:
        with zipfile.ZipFile(path) as archive:
            invalid = archive.testzip()
            if invalid:
                raise CalculationError(f"Excel 壓縮內容損壞：{invalid}")
    except zipfile.BadZipFile as exc:
        raise CalculationError("上傳檔案不是有效 Excel 活頁簿。") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    required_sheets = {"固定宣告利率", "GP", "SET", "大表 (2)", "利益分析表_基本"}
    missing_sheets = sorted(required_sheets.difference(workbook.sheetnames))
    if missing_sheets:
        raise CalculationError("建議書缺少必要工作表：" + "、".join(missing_sheets))

    name = str(read_named_value(workbook, "商品名稱")).strip()
    source_code = str(read_named_value(workbook, "險別")).strip().upper()
    if not name or not source_code:
        raise CalculationError("商品名稱或險別不可為空白。")
    currency = currency_from_value(read_named_value(workbook, "幣別"))
    declared_rate = as_number(workbook["固定宣告利率"]["A2"].value, "固定宣告利率")
    scheduled_rate = as_number(read_named_value(workbook, "預定利率"), "預定利率")
    unit_size = as_number(read_named_value(workbook, "保額單位"), "保額單位")
    source_face_units = as_number(read_named_value(workbook, "保額"), "保額")
    default_age = int(as_number(read_named_value(workbook, "保險年齡"), "保險年齡"))
    source_gender = int(as_number(read_named_value(workbook, "性別"), "性別"))
    if source_gender not in {1, 2}:
        raise CalculationError("性別欄位必須為 1（男）或 2（女）。")

    payment_mode = int(as_number(read_named_value(workbook, "繳別"), "繳別"))
    rate_by_gender: dict[str, dict[str, float]] = {"male": {}, "female": {}}
    for row in workbook["GP"].iter_rows(values_only=True):
        if len(row) < 10:
            continue
        row_code = str(row[1] or "").strip().upper()
        if row_code != source_code:
            continue
        try:
            age, mode, gender, rate = int(row[4]), int(row[5]), int(row[6]), float(row[9])
        except (TypeError, ValueError):
            continue
        if mode != payment_mode or gender not in {1, 2} or rate <= 0:
            continue
        rate_by_gender["male" if gender == 1 else "female"][str(age)] = rate
    if not rate_by_gender["male"] or not rate_by_gender["female"]:
        raise CalculationError("GP 費率表找不到所選險別、繳別與男女費率。")

    discount_tiers: list[dict[str, float]] = []
    for threshold, discount in workbook["SET"].iter_rows(min_row=29, max_row=38, min_col=23, max_col=24, values_only=True):
        if threshold is None or discount is None:
            continue
        try:
            discount_tiers.append({"minTablePremium": float(threshold), "discount": float(discount)})
        except (TypeError, ValueError):
            continue
    if not discount_tiers:
        discount_tiers = [{"minTablePremium": 0.0, "discount": 0.0}]
    discount_tiers.sort(key=lambda item: item["minTablePremium"])

    premium_name = "首期實繳保費"
    default_premium = as_number(read_named_value(workbook, premium_name), premium_name)
    table_premium = read_named_value(workbook, "表定保費", required=False)
    if table_premium is None:
        table_premium = read_named_value(workbook, "表定保險費", required=False)
    reference_name = "保費融資參考之保單現金價值" if "保費融資參考之保單現金價值" in workbook.defined_names else "首日解約金"
    reference_value = as_number(read_named_value(workbook, reference_name), reference_name)
    input_cells = {
        "age": name_destination(workbook, "保險年齡"),
        "gender": name_destination(workbook, "性別"),
        "faceUnits": name_destination(workbook, "保額"),
        "declaredRate": ("固定宣告利率", "A2"),
    }
    dividend_option_input, source_dividend_option = find_dividend_option_input(path, workbook)
    input_cells["dividendOption"] = dividend_option_input
    profile_hash = sha256_file(path)
    product_key = hashlib.sha256(normalized_text(name).encode("utf-8")).hexdigest()[:16]
    return {
        "id": f"product-{product_key}",
        "name": name,
        "normalizedName": normalized_text(name),
        "sourceCode": source_code,
        "webCode": source_code,
        "shortName": short_product_name(name, source_code),
        "currency": currency,
        "declaredRate": declared_rate,
        "scheduledRate": scheduled_rate,
        "unitSize": unit_size,
        "paymentMode": payment_mode,
        "defaultAge": default_age,
        "defaultGender": "male" if source_gender == 1 else "female",
        "defaultFaceAmount": source_face_units * unit_size,
        "defaultPremium": default_premium,
        "defaultTablePremium": float(table_premium) if table_premium is not None else None,
        "defaultReferenceCashValue": reference_value,
        "dividendOption": DIVIDEND_OPTION_CONTINUE,
        "sourceDividendOption": source_dividend_option,
        "referenceName": reference_name,
        "inputCells": {key: {"sheet": sheet, "cell": cell} for key, (sheet, cell) in input_cells.items()},
        "rateByGender": rate_by_gender,
        "discountTiers": discount_tiers,
        "sourceHash": profile_hash,
        "sourceFilename": filename,
        "sourcePath": str(path.resolve()),
        "sourceStorage": "workspace",
        "importedAt": utc_now(),
        "version": 1,
    }


def patch_sheet_xml(raw: bytes, updates: dict[str, Any]) -> bytes:
    root = ET.fromstring(raw)
    for cell_ref, value in updates.items():
        cell = root.find(f".//{{{NS_MAIN}}}c[@r='{cell_ref}']")
        if cell is None:
            raise CalculationError(f"無法定位建議書輸入儲存格 {cell_ref}。")
        for child in list(cell):
            if child.tag in {f"{{{NS_MAIN}}}f", f"{{{NS_MAIN}}}v", f"{{{NS_MAIN}}}is"}:
                cell.remove(child)
        if isinstance(value, str):
            cell.attrib["t"] = "inlineStr"
            inline = ET.SubElement(cell, f"{{{NS_MAIN}}}is")
            text = ET.SubElement(inline, f"{{{NS_MAIN}}}t")
            text.text = value
        else:
            cell.attrib.pop("t", None)
            value_node = ET.SubElement(cell, f"{{{NS_MAIN}}}v")
            value_node.text = format(float(value), ".15g")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def force_recalculation_xml(raw: bytes) -> bytes:
    root = ET.fromstring(raw)
    calc = root.find(f"{{{NS_MAIN}}}calcPr")
    if calc is None:
        calc = ET.SubElement(root, f"{{{NS_MAIN}}}calcPr")
    calc.attrib.update({"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def clear_formula_caches_xml(raw: bytes) -> bytes:
    """Remove stale formula caches so Calc must evaluate the changed inputs."""
    root = ET.fromstring(raw)
    for cell in root.findall(f".//{{{NS_MAIN}}}c"):
        if cell.find(f"{{{NS_MAIN}}}f") is None:
            continue
        cached = cell.find(f"{{{NS_MAIN}}}v")
        if cached is not None:
            cell.remove(cached)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def create_quote_workbook(source: Path, destination: Path, profile: dict[str, Any], age: int, gender: str, face_amount: float, declared_rate: float) -> None:
    unit_size = float(profile["unitSize"])
    face_units = round(face_amount / unit_size)
    if face_units <= 0:
        raise CalculationError("保額必須大於 0。")
    target_by_sheet: dict[str, dict[str, Any]] = {}
    updates = {
        "age": age,
        "gender": 2 if gender == "female" else 1,
        "faceUnits": face_units,
        "declaredRate": declared_rate,
        "dividendOption": DIVIDEND_OPTION_CONTINUE,
    }
    for key, value in updates.items():
        location = profile["inputCells"][key]
        target_by_sheet.setdefault(location["sheet"], {})[location["cell"]] = value
    sheet_paths = workbook_sheet_paths(source)
    with zipfile.ZipFile(source, "r") as reader, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as writer:
        for item in reader.infolist():
            if item.filename == "xl/calcChain.xml":
                continue
            data = reader.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = force_recalculation_xml(data)
            elif item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = clear_formula_caches_xml(data)
            for sheet_name, sheet_updates in target_by_sheet.items():
                if sheet_paths.get(sheet_name) == item.filename:
                    data = patch_sheet_xml(data, sheet_updates)
                    break
            writer.writestr(item, data)


def recalculate_with_libreoffice(source: Path, workspace: Path) -> Path:
    out_dir = workspace / "out"
    profile_dir = workspace / "profile"
    out_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    command = [
        SOFFICE,
        "--headless",
        f"-env:UserInstallation={profile_dir.as_uri()}",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(out_dir),
        str(source),
    ]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=CALCULATION_TIMEOUT_SECONDS,
            check=False,
        )
    except FileNotFoundError as exc:
        raise CalculationError("找不到 LibreOffice。請確認 soffice 已安裝並設定 SOFFICE_PATH。") from exc
    except subprocess.TimeoutExpired as exc:
        raise CalculationError(f"LibreOffice 重算逾時（{CALCULATION_TIMEOUT_SECONDS} 秒）。") from exc
    result = out_dir / source.name
    if completed.returncode != 0 or not result.exists():
        details = (completed.stderr or completed.stdout or "未知錯誤").strip()
        raise CalculationError(f"LibreOffice 無法重算建議書：{details[:500]}")
    return result


def profile_source_path(profile: dict[str, Any]) -> Path:
    storage = profile.get("sourceStorage")
    if storage == "imported":
        path = SOURCES_DIR / profile["storedFilename"]
    else:
        path = ROOT / profile["sourceFilename"]
    if not path.exists():
        raise CalculationError(f"商品來源檔不存在：{profile.get('sourceFilename', '未知檔案')}")
    return path


CALCULATION_LOCK = threading.Lock()


def calculate_quote(profile: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
    age = int(as_number(request.get("age"), "保險年齡"))
    gender = "female" if request.get("gender") == "female" else "male"
    face_amount = as_number(request.get("faceAmount"), "保額")
    declared_rate = as_number(request.get("declaredRate"), "宣告利率") / 100
    if not 0 <= declared_rate <= MAX_DECLARED_RATE:
        raise CalculationError("宣告利率必須介於 0% 與 20% 之間。")
    available_ages = sorted(int(item) for item in profile["rateByGender"][gender])
    if age not in available_ages:
        raise CalculationError(f"此商品的{('女性' if gender == 'female' else '男性')}費率表不支援 {age} 歲。")
    unit_size = float(profile["unitSize"])
    face_units = round(face_amount / unit_size)
    if abs(face_amount - face_units * unit_size) > 0.0001:
        raise CalculationError(f"保額須為 {unit_size:,.0f} {profile['currency']} 的整數倍。")

    with CALCULATION_LOCK, tempfile.TemporaryDirectory(prefix="insurance-quote-") as tmp:
        workspace = Path(tmp)
        source = profile_source_path(profile)
        input_copy = workspace / "proposal.xlsx"
        create_quote_workbook(source, input_copy, profile, age, gender, face_amount, declared_rate)
        result_path = recalculate_with_libreoffice(input_copy, workspace)
        workbook = openpyxl.load_workbook(result_path, data_only=True, read_only=True)
        benefit = numeric_year_values(read_standard_result_column(workbook, "R"), "總保障")
        cash_value = numeric_year_values(read_standard_result_column(workbook, "O"), "總保單現金價值")
        premium = as_number(read_named_value(workbook, "首期實繳保費"), "首期實繳保費")
        table_premium = read_named_value(workbook, "表定保費", required=False)
        if table_premium is None:
            table_premium = read_named_value(workbook, "表定保險費", required=False)
        reference_cash_value = as_number(read_named_value(workbook, profile["referenceName"]), profile["referenceName"])

    return {
        "profileId": profile["id"],
        "webCode": profile["webCode"],
        "currency": profile["currency"],
        "age": age,
        "gender": gender,
        "faceAmount": face_units * unit_size,
        "declaredRate": round(declared_rate * 100, 4),
        "premium": premium,
        "tablePremium": float(table_premium) if table_premium is not None else None,
        "referenceCashValue": reference_cash_value,
        "dividendOption": profile["dividendOption"],
        "benefitByYear": benefit,
        "cashValueByYear": cash_value,
        "calculationEngine": "LibreOffice Calc",
        "sourceHash": profile["sourceHash"],
    }


def build_default_quote_cache(profile: dict[str, Any]) -> dict[str, Any]:
    """Calculate the exact source-workbook scenario used by fast mode.

    The browser may only use this response when every source condition matches.
    Keeping the source hash and the complete Calc output together prevents a
    newly declared rate from accidentally reusing an older static table.
    """
    quote = calculate_quote(profile, {
        "profileId": profile["id"],
        "age": profile["defaultAge"],
        "gender": profile["defaultGender"],
        "faceAmount": profile["defaultFaceAmount"],
        "declaredRate": float(profile["declaredRate"]) * 100,
    })
    fields = (
        "sourceHash", "webCode", "currency", "age", "gender", "faceAmount",
        "declaredRate", "premium", "tablePremium", "referenceCashValue",
        "dividendOption", "benefitByYear", "cashValueByYear",
    )
    return {key: quote[key] for key in fields}


class ProfileStore:
    def __init__(self) -> None:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        SOURCES_DIR.mkdir(parents=True, exist_ok=True)
        PUBLISHED_QUOTES_DIR.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._profiles = self._load()
        self._manifest_marker = self._current_manifest_marker()

    @staticmethod
    def _current_manifest_marker() -> tuple[int, int] | None:
        try:
            metadata = MANIFEST_PATH.stat()
            return metadata.st_mtime_ns, metadata.st_size
        except OSError:
            return None

    def _load(self) -> list[dict[str, Any]]:
        if not MANIFEST_PATH.exists():
            return []
        try:
            data = json.loads(MANIFEST_PATH.read_text("utf-8"))
            return data.get("products", []) if isinstance(data, dict) else []
        except (OSError, json.JSONDecodeError):
            return []

    def _save(self) -> None:
        # A unique temporary name avoids cross-instance collisions on the
        # Cloud Storage FUSE mount.  Manifest writes are reserved for product
        # imports; ordinary quotation cache writes never touch this file.
        temporary = MANIFEST_PATH.with_name(f"{MANIFEST_PATH.name}.{uuid.uuid4().hex}.tmp")
        temporary.write_text(json.dumps({"schemaVersion": 1, "products": self._profiles}, ensure_ascii=False, indent=2) + "\n", "utf-8")
        os.replace(temporary, MANIFEST_PATH)
        self._manifest_marker = self._current_manifest_marker()

    def refresh_if_changed(self) -> None:
        """Pick up an administrator's import made by another Cloud Run instance."""
        marker = self._current_manifest_marker()
        if marker == self._manifest_marker:
            return
        with self._lock:
            # Check again after taking the local lock so quote writes do not
            # replace a freshly loaded manifest in this process.
            marker = self._current_manifest_marker()
            if marker == self._manifest_marker:
                return
            latest = self._load()
            # If an import completed while we were reading, retry once so the
            # marker and parsed JSON are from the same completed revision.
            after_read = self._current_manifest_marker()
            if after_read != marker:
                latest = self._load()
                after_read = self._current_manifest_marker()
            self._profiles = latest
            self._manifest_marker = after_read

    def _seed_default_quote_cache(self, profile: dict[str, Any]) -> bool:
        """Make every source-default Calc result available as a browser data file."""
        quote = profile.get("defaultQuote")
        if not isinstance(quote, dict) or quote.get("sourceHash") != profile.get("sourceHash"):
            return False
        cache = profile.setdefault("quoteCache", {})
        key = quote_cache_key(quote)
        existing = cache.get(key)
        if isinstance(existing, dict) and existing.get("quote") == quote:
            return False
        cache[key] = {"quote": copy.deepcopy(quote), "cachedAt": utc_now(), "kind": "source-default"}
        return True

    def _publish_profile_locked(self, profile: dict[str, Any]) -> None:
        """Atomically publish the non-sensitive browser JS data for one product."""
        PUBLISHED_QUOTES_DIR.mkdir(parents=True, exist_ok=True)
        target = PUBLISHED_QUOTES_DIR / f"{profile['id']}.js"
        temporary = target.with_name(f"{target.name}.{uuid.uuid4().hex}.tmp")
        temporary.write_text(quote_data_javascript(profile), "utf-8")
        os.replace(temporary, target)

    def publish_all(self) -> None:
        with self._lock:
            changed = False
            for profile in self._profiles:
                changed = self._seed_default_quote_cache(profile) or changed
                self._publish_profile_locked(profile)
            if changed:
                self._save()

    def bootstrap_workspace_sources(self) -> None:
        changed = False
        for path in sorted(ROOT.glob("*.xlsx")):
            if not has_financing_filename(path.name):
                continue
            try:
                profile = extract_profile(path)
                _, action = self.upsert(profile, persist=False)
                changed = changed or action != "unchanged"
            except CalculationError:
                continue
        if changed:
            self._save()
        self.publish_all()

    def upsert(self, profile: dict[str, Any], persist: bool = True) -> tuple[dict[str, Any], str]:
        # Keep the manifest independent from a caller that reuses and edits its
        # extraction dictionary for a subsequent same-name upload.
        profile = copy.deepcopy(profile)
        self.refresh_if_changed()
        with self._lock:
            matches = [item for item in self._profiles if item.get("normalizedName") == profile["normalizedName"]]
            if matches:
                existing = matches[0]
                if existing.get("sourceHash") == profile.get("sourceHash"):
                    # Older manifests predate the mandatory dividend-option
                    # control. Enrich them during bootstrap without treating an
                    # unchanged source file as a new product revision.
                    enriched = False
                    if "dividendOption" not in existing:
                        existing["dividendOption"] = profile["dividendOption"]
                        enriched = True
                    if "sourceDividendOption" not in existing:
                        existing["sourceDividendOption"] = profile["sourceDividendOption"]
                        enriched = True
                    existing_inputs = existing.setdefault("inputCells", {})
                    if "dividendOption" not in existing_inputs:
                        existing_inputs["dividendOption"] = profile["inputCells"]["dividendOption"]
                        enriched = True
                    if profile.get("defaultQuote") and existing.get("defaultQuote") != profile["defaultQuote"]:
                        existing["defaultQuote"] = profile["defaultQuote"]
                        enriched = True
                    if self._seed_default_quote_cache(existing):
                        enriched = True
                    if enriched and persist:
                        self._save()
                        self._publish_profile_locked(existing)
                    return existing, "unchanged"
                profile["id"] = existing["id"]
                profile["webCode"] = existing.get("webCode") or profile["sourceCode"]
                profile["version"] = int(existing.get("version", 1)) + 1
                index = self._profiles.index(existing)
                self._profiles[index] = profile
                action = "updated"
            else:
                used_codes = {item.get("webCode") for item in self._profiles}
                if profile["sourceCode"] in used_codes:
                    profile["webCode"] = f"IMP-{profile['id'].split('-')[-1][:8].upper()}"
                profile["version"] = int(profile.get("version", 1))
                self._profiles.append(profile)
                action = "created"
            self._seed_default_quote_cache(profile)
            if persist:
                self._save()
            self._publish_profile_locked(profile)
            return profile, action

    def cache_quote(self, profile_id: str, quote: dict[str, Any]) -> None:
        """Persist an exact Calc result so all later visitors can use it instantly."""
        with self._lock:
            profile = next((item for item in self._profiles if item.get("id") == profile_id), None)
            if not profile or quote.get("sourceHash") != profile.get("sourceHash"):
                return
            # Keep user-requested results outside products.json.  With this
            # design, independent Cloud Run instances can publish different
            # quotes in parallel without a read-modify-write race.
            target = quote_record_path(profile_id, quote)
            target.parent.mkdir(parents=True, exist_ok=True)
            temporary = target.with_name(f"{target.name}.{uuid.uuid4().hex}.tmp")
            temporary.write_text(json.dumps({
                "schemaVersion": 1,
                "cacheKey": quote_cache_key(quote),
                "cachedAt": utc_now(),
                "quote": quote,
            }, ensure_ascii=False, separators=(",", ":")), "utf-8")
            os.replace(temporary, target)
            self._publish_profile_locked(profile)

    def import_file(self, stream: io.BufferedReader, filename: str) -> tuple[dict[str, Any], str]:
        filename = safe_filename(filename)
        if not has_financing_filename(filename):
            raise CalculationError("只接受檔名包含「保費融資」的建議書。")
        suffix = Path(filename).suffix.lower()
        if suffix not in {".xlsx", ".xlsm"}:
            raise CalculationError("只接受 .xlsx 或 .xlsm 建議書。")
        with tempfile.TemporaryDirectory(prefix="insurance-import-") as tmp:
            temporary = Path(tmp) / f"upload{suffix}"
            written = 0
            with temporary.open("wb") as output:
                while True:
                    chunk = stream.read(1024 * 1024)
                    if not chunk:
                        break
                    written += len(chunk)
                    if written > MAX_UPLOAD_BYTES:
                        raise CalculationError("建議書超過 40 MB 上限。")
                    output.write(chunk)
            if written == 0:
                raise CalculationError("上傳檔案為空白。")
            profile = extract_profile(temporary, filename)
            stored_filename = f"{profile['sourceHash']}{suffix}"
            target = SOURCES_DIR / stored_filename
            shutil.copy2(temporary, target)
        profile["sourceStorage"] = "imported"
        profile["storedFilename"] = stored_filename
        profile.pop("sourcePath", None)
        try:
            profile["defaultQuote"] = build_default_quote_cache(profile)
        except Exception:
            # A proposal is not accepted unless its source-default fast quote is
            # independently recalculated.  Remove its staged copy on failure.
            target.unlink(missing_ok=True)
            raise
        return self.upsert(profile)

    def get(self, profile_id: str) -> dict[str, Any] | None:
        self.refresh_if_changed()
        with self._lock:
            return next((item for item in self._profiles if item.get("id") == profile_id), None)

    def public_profiles(self) -> list[dict[str, Any]]:
        fields = {
            "id", "name", "sourceCode", "webCode", "shortName", "currency", "declaredRate", "scheduledRate", "unitSize",
            "defaultAge", "defaultGender", "defaultFaceAmount", "defaultPremium", "version", "sourceFilename", "sourceHash",
            "rateByGender", "discountTiers", "dividendOption", "defaultQuote",
        }
        self.refresh_if_changed()
        with self._lock:
            public: list[dict[str, Any]] = []
            for item in self._profiles:
                profile = {key: value for key, value in item.items() if key in fields}
                profile["quoteDataPath"] = f"/api/product-data/{item['id']}.js"
                profile["quoteDataVersion"] = quote_data_version(item)
                public.append(profile)
            return public


STORE = ProfileStore()


class RequestHandler(SimpleHTTPRequestHandler):
    server_version = "InsuranceCalculation/1.0"

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def end_headers(self) -> None:
        origin = self.headers.get("Origin", "").rstrip("/")
        if origin and origin in CORS_ALLOWED_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type, X-Insurance-Admin-Token")
        super().end_headers()

    def send_json(self, payload: dict[str, Any], status: HTTPStatus = HTTPStatus.OK) -> None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def send_javascript(self, contents: str, version: str) -> None:
        data = contents.encode("utf-8")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/javascript; charset=utf-8")
        self.send_header("Cache-Control", "public, max-age=31536000, immutable")
        self.send_header("ETag", f'"{version}"')
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_OPTIONS(self) -> None:  # noqa: N802
        self.send_response(HTTPStatus.NO_CONTENT)
        self.end_headers()

    def do_GET(self) -> None:  # noqa: N802
        route = urlparse(self.path).path
        if route == "/api/health":
            self.send_json({"ok": True, "engine": "LibreOffice Calc"})
            return
        if route == "/api/products":
            self.send_json({"products": STORE.public_profiles()})
            return
        match = re.fullmatch(r"/api/product-data/([A-Za-z0-9-]+)\.js", route)
        if match:
            profile = STORE.get(match.group(1))
            if not profile:
                self.send_json({"error": "找不到商品資料檔。"}, HTTPStatus.NOT_FOUND)
                return
            self.send_javascript(quote_data_javascript(profile), quote_data_version(profile))
            return
        super().do_GET()

    def do_POST(self) -> None:  # noqa: N802
        route = urlparse(self.path).path
        try:
            if route == "/api/quote":
                self.handle_quote()
                return
            if route == "/api/import":
                self.handle_import()
                return
            self.send_json({"error": "找不到 API 路徑。"}, HTTPStatus.NOT_FOUND)
        except PermissionError as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.UNAUTHORIZED)
        except CalculationError as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.UNPROCESSABLE_ENTITY)
        except Exception as exc:  # pragma: no cover - protects the local HTTP boundary
            self.log_error("Unexpected error: %s", exc)
            self.send_json({"error": "本機試算服務發生未預期錯誤。"}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def handle_quote(self) -> None:
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0 or length > 64 * 1024:
            raise CalculationError("試算要求大小無效。")
        payload = json.loads(self.rfile.read(length).decode("utf-8"))
        if not isinstance(payload, dict):
            raise CalculationError("試算要求格式無效。")
        profile_id = str(payload.get("profileId", ""))
        profile = STORE.get(profile_id)
        if not profile:
            raise CalculationError("找不到已匯入的商品來源。")
        quote = calculate_quote(profile, payload)
        STORE.cache_quote(profile_id, quote)
        self.send_json({"quote": quote})

    def handle_import(self) -> None:
        if ADMIN_UPLOAD_TOKEN:
            provided = self.headers.get("X-Insurance-Admin-Token", "")
            if not hmac.compare_digest(provided, ADMIN_UPLOAD_TOKEN):
                raise PermissionError("未授權匯入。請輸入管理者上傳金鑰。")
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            raise CalculationError("匯入要求必須使用 multipart/form-data。")
        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": content_type},
        )
        if "proposal" not in form:
            raise CalculationError("請選擇建議書檔案。")
        uploaded = form["proposal"]
        if not getattr(uploaded, "file", None) or not getattr(uploaded, "filename", None):
            raise CalculationError("請選擇建議書檔案。")
        profile, action = STORE.import_file(uploaded.file, uploaded.filename)
        self.send_json({"action": action, "product": STORE.public_profiles()[[item["id"] for item in STORE.public_profiles()].index(profile["id"])]})

    def log_message(self, format: str, *args: Any) -> None:
        print(f"[{self.log_date_time_string()}] {format % args}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Serve the premium-financing calculator with LibreOffice-backed quotations.")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()
    STORE.bootstrap_workspace_sources()
    STORE.publish_all()
    server = ThreadingHTTPServer((args.host, args.port), RequestHandler)
    print(f"Insurance calculator: http://{args.host}:{args.port}/index.html")
    print("Calculation engine: LibreOffice Calc")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
